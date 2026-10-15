"""
Справочное окно «Пересчёт мощностей радиаторов LaggarTT».

Показывает полную таблицу мощностей (7 типов × 27 длин) для выбранной высоты
при текущем температурном графике.

Структура окна:
    • Фиксированная шапка сверху: заголовок + кнопки «Экспорт в Excel» и «Закрыть»
    • Прокручиваемая область ниже: график, выбор высоты, таблица,
      примечание, 4 информационных блока

Адаптивность:
    • Ширины колонок рассчитываются от ширины окна
    • Размер шрифта в таблице подбирается так, чтобы текст влезал в ячейку
    • При увеличении масштаба шрифт в таблице уменьшается, но не ниже 6 pt

Возможности:
    • переключение высоты радиокнопками (300/400/500/600/900)
    • мгновенный пересчёт таблицы
    • экспорт всех 5 высот в один Excel-файл (5 листов + лист «Методика»)
"""

import os
import platform
import subprocess
from datetime import datetime
from typing import Dict, Optional

import tkinter as tk
import tkinter.font as tkfont
from tkinter import ttk, messagebox

from nextt.logger import get_logger
from nextt.events import bus, Events
from nextt.core.power_recalculator import (
    PowerRecalculator,
    EXPONENTS,
    VALID_TYPES,
    VALID_LENGTHS,
    VALID_HEIGHTS,
)
from nextt.utils.window_manager import get_safe_geometry, setup_dialog_window

logger = get_logger(__name__)


# =========================================================================
# Цвета
# =========================================================================
COLOR_BG = "#f5f5f5"
COLOR_CARD_BG = "#fafafa"
COLOR_CARD_BORDER = "#c8c8c8"
COLOR_HEADING = "#263168"
COLOR_TEXT = "#2c3e50"
COLOR_SUBTEXT = "#555555"
COLOR_FORMULA_BG = "#f0f4ff"
COLOR_FORMULA_FG = "#1a2456"

MONO_FAMILY = "Consolas"

# Базовые ширины колонок (при scale = 1.0)
BASE_CELL_WIDTH = 40
BASE_TYPE_WIDTH = 80

# Минимальный размер шрифта в таблице мощностей
MIN_TABLE_FONT_SIZE = 6

# Базовый размер шрифта в таблице (при scale = 1.0)
BASE_TABLE_FONT_SIZE = 9


class PowerReferenceDialog:
    """
    Справочное окно мощностей.

    Показывает таблицу мощностей для выбранной высоты и текущего графика.
    НЕ влияет на спецификацию и НЕ изменяет данные приложения.
    """

    def __init__(self, app, graph: Dict[str, int], delta_t: float):
        self.app = app
        self.font_manager = app.font_manager

        self._graph = {
            "supply": int(graph.get("supply", 95)),
            "return": int(graph.get("return", 85)),
            "room":   int(graph.get("room",   20)),
        }
        self._delta_t = float(delta_t)

        self._selected_height = tk.IntVar(value=500)

        self._dialog: Optional[tk.Toplevel] = None
        self._tree: Optional[ttk.Treeview] = None
        self._table_frame: Optional[tk.Frame] = None
        self._canvas: Optional[tk.Canvas] = None
        self._content_frame: Optional[tk.Frame] = None

        # Флаг: заблокировать пересчёт ширин во время первого построения
        self._layout_ready = False

        # Текущий подобранный шрифт таблицы (для переиспользования)
        self._table_font: Optional[tkfont.Font] = None
        self._table_font_bold: Optional[tkfont.Font] = None

    # ==================================================================
    # Публичный метод
    # ==================================================================

    def show(self) -> None:
        """Открывает диалог."""
        if self._dialog and self._dialog.winfo_exists():
            self._dialog.lift()
            self._dialog.focus_force()
            return

        self._create_dialog()

        # Подписка на изменение масштаба шрифтов
        bus.subscribe(Events.FONT_SCALE_CHANGED, self._on_font_scale_changed)

        # Заполняем таблицу и рассчитываем ширины после отрисовки окна
        self._dialog.after(80, self._first_layout)

    def _first_layout(self) -> None:
        """Первичное построение: рассчитываем ширины, подбираем шрифт, заполняем таблицу."""
        if not self._dialog or not self._dialog.winfo_exists():
            return
        self._recalc_column_widths()
        self._fill_tree()
        self._layout_ready = True

    # ==================================================================
    # Шрифты
    # ==================================================================

    def _font_default(self):
        return self.font_manager.get_default_font()

    def _font_bold(self, base_size: int = 9):
        return self.font_manager.get_bold_font(base_size)

    def _font_heading(self):
        return self.font_manager.get_heading_font()

    def _font_title(self):
        return self.font_manager.get_title_font()

    def _font_mono(self, base_size: int = 10):
        try:
            scale = self.font_manager.scale
        except Exception:
            scale = 1.0
        size = max(8, int(base_size * scale))
        return (MONO_FAMILY, size)

    # ==================================================================
    # Создание окна
    # ==================================================================

    def _create_dialog(self) -> None:
        self._dialog = tk.Toplevel(self.app.root)

        setup_dialog_window(self._dialog, "Пересчёт мощностей радиаторов LaggarTT")

        try:
            from nextt.config import get_resource_path
            icon_path = get_resource_path("icon.ico")
            if icon_path and os.path.exists(icon_path):
                self._dialog.iconbitmap(icon_path)
        except Exception as e:
            logger.debug(f"Не удалось загрузить иконку: {e}")

        self._dialog.geometry(get_safe_geometry(self._dialog))
        self._dialog.configure(bg=COLOR_BG)

        default_font = self._font_default()
        heading_font = self._font_heading()

        # --- Главный контейнер ---
        outer = tk.Frame(self._dialog, bg=COLOR_BG)
        outer.pack(fill="both", expand=True)

        # ================================================================
        # ФИКСИРОВАННАЯ ШАПКА: заголовок слева, кнопки справа
        # ================================================================
        top_bar = tk.Frame(outer, bg=COLOR_BG)
        top_bar.pack(fill="x", padx=20, pady=(14, 10))

        tk.Label(
            top_bar,
            text="📊 МОЩНОСТИ РАДИАТОРОВ ПРИ ТЕКУЩЕМ ГРАФИКЕ",
            font=heading_font,
            fg=COLOR_HEADING,
            bg=COLOR_BG,
            anchor="w",
        ).pack(side="left", anchor="w")

        btn_close = ttk.Button(
            top_bar,
            text="Закрыть",
            command=self._close,
        )
        btn_close.pack(side="right", padx=(6, 0))

        btn_export = ttk.Button(
            top_bar,
            text="Экспорт в Excel",
            command=self._on_export,
        )
        btn_export.pack(side="right")

        # ================================================================
        # ПРОКРУЧИВАЕМАЯ ОБЛАСТЬ
        # ================================================================
        canvas_container = tk.Frame(outer, bg=COLOR_BG)
        canvas_container.pack(fill="both", expand=True)

        self._canvas = tk.Canvas(
            canvas_container,
            bg=COLOR_BG,
            highlightthickness=0,
            bd=0,
        )
        vsb = ttk.Scrollbar(
            canvas_container,
            orient="vertical",
            command=self._canvas.yview,
        )
        self._canvas.configure(yscrollcommand=vsb.set)

        self._canvas.pack(side="left", fill="both", expand=True)
        vsb.pack(side="right", fill="y")

        self._content_frame = tk.Frame(self._canvas, bg=COLOR_BG)
        self._content_window = self._canvas.create_window(
            (0, 0),
            window=self._content_frame,
            anchor="nw",
        )

        self._content_frame.bind(
            "<Configure>",
            lambda e: self._canvas.configure(scrollregion=self._canvas.bbox("all")),
        )

        self._canvas.bind("<Configure>", self._on_canvas_configure)

        self._dialog.bind("<MouseWheel>", self._on_mousewheel)
        self._canvas.bind("<MouseWheel>", self._on_mousewheel)
        self._content_frame.bind("<MouseWheel>", self._on_mousewheel)

        # ================================================================
        # СОДЕРЖИМОЕ
        # ================================================================
        content = self._content_frame

        # --- Информация о графике ---
        info = tk.Frame(content, bg=COLOR_BG)
        info.pack(fill="x", padx=20, pady=(0, 10))

        graph_str = f"{self._graph['supply']}/{self._graph['return']}/{self._graph['room']}"
        delta_str = f"{self._delta_t:.2f}".replace('.', ',')

        tk.Label(
            info,
            text=f"Температурный график: {graph_str} °С   (Δt = {delta_str} °С)",
            font=self._font_bold(9),
            fg=COLOR_HEADING,
            bg=COLOR_BG,
            anchor="w",
        ).pack(anchor="w")

        tk.Label(
            info,
            text="Методика: ГОСТ 31311-2022.  Формула: Q = Qном × (Δt/70)ⁿ",
            font=default_font,
            fg=COLOR_TEXT,
            bg=COLOR_BG,
            anchor="w",
        ).pack(anchor="w", pady=(4, 0))

        # --- Блок выбора высоты ---
        height_frame = ttk.LabelFrame(content, text="ВЫСОТА РАДИАТОРА", padding=8)
        height_frame.pack(fill="x", padx=20, pady=(0, 10))

        for h in VALID_HEIGHTS:
            ttk.Radiobutton(
                height_frame,
                text=f"{h} мм",
                variable=self._selected_height,
                value=h,
                command=self._on_height_changed
            ).pack(side="left", padx=10)

        # --- Таблица ---
        self._table_frame = tk.Frame(content, bg=COLOR_BG)
        self._table_frame.pack(fill="x", padx=20, pady=(0, 10))

        columns = ["type"] + [str(L) for L in VALID_LENGTHS]
        self._tree = ttk.Treeview(
            self._table_frame,
            columns=columns,
            show="headings",
            height=7,
            selectmode="none",
        )

        # Заголовки и начальные ширины
        self._tree.heading("type", text="Тип \\ Длина")
        self._tree.column("type", width=BASE_TYPE_WIDTH, anchor="center",
                          minwidth=BASE_TYPE_WIDTH, stretch=False)

        for L in VALID_LENGTHS:
            self._tree.heading(str(L), text=str(L))
            self._tree.column(str(L), width=BASE_CELL_WIDTH, anchor="center",
                              minwidth=BASE_CELL_WIDTH, stretch=False)

        # Теги для строк (шрифт будет назначен в _fit_table_font)
        self._tree.tag_configure("even", background="#FFFFFF")
        self._tree.tag_configure("odd", background="#F0F0F0")

        self._tree.pack(fill="x", expand=False)

        self._table_frame.bind("<Configure>", self._on_table_resize)

        # --- Примечание ---
        tk.Label(
            content,
            text="ℹ️  При Δt = 70 °С (95/85/20) значения совпадают с паспортными. "
                 "Значения округлены до целых ватт.",
            font=default_font,
            fg=COLOR_SUBTEXT,
            bg=COLOR_BG,
            anchor="w",
        ).pack(anchor="w", padx=20, pady=(0, 10))

        # --- Информационные блоки 2×2 ---
        self._build_info_sections(content)

    def _on_canvas_configure(self, event) -> None:
        """Растягивает содержимое по ширине canvas."""
        if not self._canvas or not self._content_frame:
            return
        self._canvas.itemconfigure(self._content_window, width=event.width)

    def _on_mousewheel(self, event) -> None:
        """Прокрутка колёсиком мыши."""
        if self._canvas:
            self._canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")

    # ==================================================================
    # Адаптивные ширины колонок + подбор шрифта
    # ==================================================================

    def _on_table_resize(self, event) -> None:
        """Пересчёт ширин при изменении размера таблицы."""
        if not self._layout_ready:
            return
        self._recalc_column_widths()

    def _recalc_column_widths(self) -> None:
        """Рассчитывает ширины колонок и подбирает шрифт таблицы."""
        if not self._tree or not self._table_frame:
            return
        if not self._table_frame.winfo_exists():
            return

        self._table_frame.update_idletasks()
        available = self._table_frame.winfo_width()

        if available < 100:
            available = self._dialog.winfo_screenwidth() - 100

        # Небольшой запас на рамки
        available -= 6
        if available < 200:
            available = 200

        num_cols = 1 + len(VALID_LENGTHS)  # 1 + 27 = 28

        scale = 1.0
        try:
            scale = float(self.font_manager.scale)
        except Exception:
            scale = 1.0

        # Доп. ширина для колонки «Тип \ Длина»
        extra_for_type = int(40 * scale)

        # Вычитаем доп. ширину ДО расчёта cell_width, чтобы всё влезло
        available_for_cells = available - extra_for_type
        if available_for_cells < num_cols * 20:
            available_for_cells = num_cols * 20

        cell_width = available_for_cells // num_cols
        type_width = cell_width + extra_for_type

        # Применяем ширины
        try:
            self._tree.column("type", width=type_width, minwidth=type_width, stretch=False)
            for L in VALID_LENGTHS:
                self._tree.column(str(L), width=cell_width, minwidth=cell_width, stretch=False)
        except tk.TclError:
            pass

        # Подбираем шрифт под ширины
        self._fit_table_font(cell_width, type_width, scale)

    def _fit_table_font(self, cell_width: int, type_width: int, scale: float) -> None:
        """
        Подбирает размер шрифта для таблицы так, чтобы самый длинный текст
        («3000», «Тип \\ Длина», «Тип 10») помещался в отведённую ширину.
        Размер уменьшается от базового до MIN_TABLE_FONT_SIZE.
        """
        # Определяем семейство шрифта из FontManager
        family = "Segoe UI"
        try:
            base_font_tuple = self.font_manager.get_default_font()
            if base_font_tuple and len(base_font_tuple) >= 1:
                family = base_font_tuple[0]
        except Exception:
            pass

        # Базовый размер при scale = 1.0
        base_size = BASE_TABLE_FONT_SIZE

        # Отступ внутри ячейки (слева + справа + рамки Treeview)
        padding_cell = 10
        padding_type = 16

        # Подбираем размер от базового вниз до минимального
        chosen_size = MIN_TABLE_FONT_SIZE
        for size in range(base_size, MIN_TABLE_FONT_SIZE - 1, -1):
            try:
                test_font = tkfont.Font(family=family, size=size)
            except Exception:
                continue

            # Самый длинный текст в ячейке данных — «3000» (или «Тип 10» для колонки типа)
            w_number = test_font.measure("3000") + padding_cell
            w_type_header = test_font.measure("Тип \\ Длина") + padding_type
            w_type_data = test_font.measure("Тип 10") + padding_type

            if w_number <= cell_width and w_type_header <= type_width and w_type_data <= type_width:
                chosen_size = size
                break
            chosen_size = size

        # Создаём финальные шрифты
        try:
            self._table_font = tkfont.Font(family=family, size=chosen_size)
            self._table_font_bold = tkfont.Font(family=family, size=chosen_size, weight="bold")
        except Exception:
            self._table_font = None
            self._table_font_bold = None
            return

        # Применяем шрифт к тегам строк
        try:
            self._tree.tag_configure("even", font=self._table_font, background="#FFFFFF")
            self._tree.tag_configure("odd", font=self._table_font, background="#F0F0F0")
        except tk.TclError:
            pass

        # Применяем шрифт к заголовкам через стиль
        try:
            style = ttk.Style()
            style.configure(
                "PowerRef.Treeview.Heading",
                font=self._table_font_bold,
            )
            style.configure(
                "PowerRef.Treeview",
                font=self._table_font,
                rowheight=max(20, int(chosen_size * 2.2)),
            )
            self._tree.configure(style="PowerRef.Treeview")
        except tk.TclError:
            pass

    # ==================================================================
    # Масштабирование шрифтов
    # ==================================================================

    def _on_font_scale_changed(self, scale: float) -> None:
        """Пересчитывает ширины и шрифт при смене масштаба."""
        if not self._dialog or not self._dialog.winfo_exists():
            return
        try:
            self.font_manager.refresh_all_windows(self._dialog)
        except Exception as e:
            logger.debug(f"Ошибка обновления шрифтов: {e}")

        self._recalc_column_widths()

    # ==================================================================
    # Информационные блоки
    # ==================================================================

    def _build_info_sections(self, parent: tk.Frame) -> None:
        grid = tk.Frame(parent, bg=COLOR_BG)
        grid.pack(fill="both", expand=True, padx=20, pady=(0, 20))

        grid.columnconfigure(0, weight=1, uniform="info")
        grid.columnconfigure(1, weight=1, uniform="info")

        self._build_card_exponents(grid, row=0, col=0)
        self._build_card_formula(grid, row=0, col=1)
        self._build_card_example(grid, row=1, col=0)
        self._build_card_graph(grid, row=1, col=1)

    def _make_card(self, parent: tk.Frame, row: int, col: int) -> tk.Frame:
        outer = tk.Frame(
            parent,
            bg=COLOR_CARD_BG,
            highlightbackground=COLOR_CARD_BORDER,
            highlightthickness=1,
            bd=0,
        )
        outer.grid(row=row, column=col, sticky="nsew", padx=6, pady=6)

        inner = tk.Frame(outer, bg=COLOR_CARD_BG)
        inner.pack(fill="both", expand=True, padx=12, pady=10)
        return inner

    def _card_heading(self, parent: tk.Frame, text: str) -> None:
        tk.Label(
            parent,
            text=text.upper(),
            font=self._font_bold(9),
            fg=COLOR_HEADING,
            bg=COLOR_CARD_BG,
            anchor="w",
            justify="left",
        ).pack(anchor="w", pady=(0, 8))

    def _card_paragraph(self, parent: tk.Frame, text: str) -> None:
        tk.Label(
            parent,
            text=text,
            font=self._font_default(),
            fg=COLOR_TEXT,
            bg=COLOR_CARD_BG,
            anchor="w",
            justify="left",
            wraplength=520,
        ).pack(anchor="w", pady=2)

    def _card_formula(self, parent: tk.Frame, text: str) -> None:
        wrap = tk.Frame(parent, bg=COLOR_FORMULA_BG)
        wrap.pack(anchor="w", pady=4, fill="x")

        tk.Label(
            wrap,
            text=text,
            font=self._font_mono(10),
            fg=COLOR_FORMULA_FG,
            bg=COLOR_FORMULA_BG,
            anchor="w",
            justify="left",
            padx=8, pady=5,
        ).pack(anchor="w")

    def _card_table(self, parent: tk.Frame, columns: list, rows: list) -> None:
        table_wrap = tk.Frame(parent, bg=COLOR_CARD_BG)
        table_wrap.pack(anchor="w", pady=4, fill="x")

        for r_idx, row in enumerate(rows):
            for c_idx, cell in enumerate(row):
                is_header = (r_idx == 0)
                bg = "#e8e8e8" if is_header else "#ffffff"
                fg = COLOR_HEADING if is_header else COLOR_TEXT
                fnt = self._font_bold(9) if is_header else self._font_default()

                tk.Label(
                    table_wrap,
                    text=str(cell),
                    font=fnt,
                    fg=fg,
                    bg=bg,
                    relief="solid",
                    borderwidth=1,
                    padx=8, pady=2,
                    anchor="center",
                ).grid(row=r_idx, column=c_idx, sticky="nsew")

    # --- Блоки ---

    def _build_card_exponents(self, parent: tk.Frame, row: int, col: int) -> None:
        inner = self._make_card(parent, row, col)

        self._card_heading(inner, "Экспоненты n по типам радиаторов")

        self._card_paragraph(
            inner,
            "Экспонент n определяет чувствительность мощности к изменению "
            "температурного напора Δt. Чем больше n, тем сильнее падает "
            "мощность при снижении Δt."
        )

        headers = ["Тип"] + VALID_TYPES
        exponents_row = ["n"]
        for rt in VALID_TYPES:
            n_val = EXPONENTS.get(500, {}).get(rt, 1.0)
            exponents_row.append(f"{n_val:.2f}".replace('.', ','))

        self._card_table(inner, headers, [headers, exponents_row])

        self._card_paragraph(
            inner,
            "Значения n одинаковы для всех высот (300–900 мм)."
            
        )

    def _build_card_formula(self, parent: tk.Frame, row: int, col: int) -> None:
        inner = self._make_card(parent, row, col)

        self._card_heading(inner, "Формула пересчёта")

        self._card_paragraph(inner, "Мощность при текущем графике:")
        self._card_formula(inner, "Q = Qном × (Δt / 70)ⁿ")

        self._card_paragraph(
            inner,
            "где  Qном — паспортная мощность при 95/85/20 (Δt = 70 °С),\n"
            "       Δt — текущий температурный напор,\n"
            "       n — экспонент типа радиатора."
        )

        self._card_paragraph(inner, "Температурный напор:")
        self._card_formula(inner, "Δt = (t подачи + t обратки) / 2 − t помещения")

        self._card_paragraph(
            inner,
            "При Δt = 70 коэффициент равен 1,0 — пересчёт не меняет "
            "паспортные значения."
        )

    def _build_card_example(self, parent: tk.Frame, row: int, col: int) -> None:
        inner = self._make_card(parent, row, col)

        self._card_heading(inner, "Пример расчёта при текущем графике")

        self._card_paragraph(
            inner,
            "Возьмём радиатор Тип 10, высота 500 мм, длина 1000 мм."
        )

        qnom = 914
        try:
            df = self.app.data_provider.radiators_df
            mask = (
                df['Наименование'].str.contains('/500/1000', regex=False, na=False) &
                (df['RadiatorType'] == 10)
            )
            matches = df[mask]
            if not matches.empty:
                qnom = int(round(float(matches.iloc[0].get('Мощность, Вт', 914))))
        except Exception:
            pass

        n = EXPONENTS.get(500, {}).get('10', 1.14)
        factor = (self._delta_t / 70.0) ** n if self._delta_t > 0 else 1.0
        q_recalc = qnom * factor if not PowerRecalculator.is_reference(self._delta_t) else qnom

        delta_str = f"{self._delta_t:.2f}".replace('.', ',')
        n_str = f"{n:.2f}".replace('.', ',')
        factor_str = f"{factor:.3f}".replace('.', ',')
        q_str = int(round(q_recalc))
        qnom_str = int(round(qnom))
        graph_str = f"{self._graph['supply']}/{self._graph['return']}/{self._graph['room']}"

        example_text = (
            f"Паспортная мощность:  Qном = {qnom_str} Вт\n"
            f"Текущий график:       {graph_str}  →  Δt = {delta_str} °С\n"
            f"Экспонент (тип 10):   n = {n_str}\n"
            f"Коэффициент:          ({self._delta_t:.0f} / 70)^{n_str} = {factor_str}\n"
            f"Мощность:             Q = {qnom_str} × {factor_str} ≈ {q_str} Вт"
        )

        wrap = tk.Frame(inner, bg=COLOR_FORMULA_BG)
        wrap.pack(anchor="w", pady=4, fill="x")

        tk.Label(
            wrap,
            text=example_text,
            font=self._font_mono(10),
            fg=COLOR_FORMULA_FG,
            bg=COLOR_FORMULA_BG,
            anchor="w",
            justify="left",
            padx=8, pady=6,
        ).pack(anchor="w")

        if not PowerRecalculator.is_reference(self._delta_t):
            diff_pct = (1 - factor) * 100 if factor < 1 else (factor - 1) * 100
            sign = "упала" if factor < 1 else "выросла"
            self._card_paragraph(
                inner,
                f"Для сравнения: при эталонном 95/85/20 мощность была бы "
                f"{qnom_str} Вт. Сейчас она {sign} на {diff_pct:.0f} %."
            )
        else:
            self._card_paragraph(
                inner,
                "Текущий график эталонный — пересчёт не требуется, "
                "значения совпадают с паспортными."
            )

    def _build_card_graph(self, parent: tk.Frame, row: int, col: int) -> None:
        inner = self._make_card(parent, row, col)

        self._card_heading(inner, "Пояснение к температурному графику")

        self._card_paragraph(
            inner,
            "Текущий график записывается тремя числами через слэш:"
        )

        graph_str = f"{self._graph['supply']} / {self._graph['return']} / {self._graph['room']}"
        self._card_formula(inner, graph_str)

        self._card_paragraph(
            inner,
            f"{self._graph['supply']} — температура подающей линии, °С\n"
            f"{self._graph['return']} — температура обратной линии, °С\n"
            f"{self._graph['room']} — температура воздуха в помещении, °С"
        )

        self._card_paragraph(inner, "Типовые графики и соответствующие Δt:")

        rows = [
            ["График", "Δt, °С"],
            ["95/85/20 (эталон)", "70,0"],
            ["90/70/20", "60,0"],
            ["85/60/20", "52,5"],
            ["75/65/20", "50,0"],
            ["55/45/20", "30,0"],
        ]
        self._card_table(inner, None, rows)

    # ==================================================================
    # Заполнение таблицы
    # ==================================================================

    def _fill_tree(self) -> None:
        if not self._tree:
            return

        for item in self._tree.get_children():
            self._tree.delete(item)

        height = self._selected_height.get()
        table = PowerRecalculator.full_table_for_height(
            height, self._delta_t, self.app.data_provider
        )

        for i, rt in enumerate(VALID_TYPES):
            row_values = [f"Тип {rt}"] + [str(v) for v in table.get(rt, [])]
            tag = "even" if i % 2 == 0 else "odd"
            self._tree.insert("", "end", values=row_values, tags=(tag,))

    # ==================================================================
    # Обработчики
    # ==================================================================

    def _on_height_changed(self) -> None:
        self._fill_tree()

    # ==================================================================
    # Экспорт в Excel
    # ==================================================================

    def _on_export(self) -> None:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
            from openpyxl.utils import get_column_letter
            from pathlib import Path

            downloads = Path.home() / "Downloads"
            if not downloads.exists():
                downloads = Path.home() / "Загрузки"
            if not downloads.exists():
                downloads = Path.home()

            nextt_folder = downloads / "NexTT"
            nextt_folder.mkdir(exist_ok=True)

            graph_str = f"{self._graph['supply']}-{self._graph['return']}-{self._graph['room']}"
            base_name = f"Мощности_{graph_str}"

            file_path = nextt_folder / f"{base_name}.xlsx"
            counter = 1
            while file_path.exists():
                file_path = nextt_folder / f"{base_name}_{counter}.xlsx"
                counter += 1

            header_font = Font(name='Calibri', size=11, bold=True)
            title_font = Font(name='Calibri', size=13, bold=True, color="263168")
            section_font = Font(name='Calibri', size=11, bold=True, color="263168")
            data_font = Font(name='Calibri', size=10)
            mono_font = Font(name='Consolas', size=10, color="1A2456")

            center = Alignment(horizontal='center', vertical='center', wrap_text=False)
            left_wrap = Alignment(horizontal='left', vertical='top', wrap_text=True)

            thin = Side(style='thin', color="C8C8C8")
            border = Border(left=thin, right=thin, top=thin, bottom=thin)

            formula_fill = PatternFill(start_color="F0F4FF", end_color="F0F4FF", fill_type="solid")
            header_fill = PatternFill(start_color="E8E8E8", end_color="E8E8E8", fill_type="solid")

            wb = Workbook()
            wb.remove(wb.active)

            all_tables = PowerRecalculator.full_table_all_heights(
                self._delta_t, self.app.data_provider
            )

            # --- 5 листов с таблицами ---
            for h in VALID_HEIGHTS:
                ws = wb.create_sheet(f"{h} мм")

                ws.append([f"Мощности радиаторов LaggarTT, высота {h} мм"])
                ws.cell(row=1, column=1).font = title_font

                delta_str = f"{self._delta_t:.2f}".replace('.', ',')
                ws.append([
                    f"Температурный график: "
                    f"{self._graph['supply']}/{self._graph['return']}/{self._graph['room']} °С  "
                    f"(Δt = {delta_str} °С)"
                ])
                ws.cell(row=2, column=1).font = data_font

                ws.append(["Методика: ГОСТ 31311-2022.  Формула: Q = Qном × (Δt/70)ⁿ"])
                ws.cell(row=3, column=1).font = data_font

                ws.append([])

                header_row = ["Тип \\ Длина"] + [str(L) for L in VALID_LENGTHS]
                ws.append(header_row)

                table = all_tables[h]
                for rt in VALID_TYPES:
                    row = [f"Тип {rt}"] + list(table.get(rt, []))
                    ws.append(row)

                header_row_idx = 5
                for col in range(1, len(header_row) + 1):
                    cell = ws.cell(row=header_row_idx, column=col)
                    cell.font = header_font
                    cell.alignment = center
                    cell.border = border
                    cell.fill = header_fill

                for r_idx in range(header_row_idx + 1, header_row_idx + 1 + len(VALID_TYPES)):
                    for col in range(1, len(header_row) + 1):
                        cell = ws.cell(row=r_idx, column=col)
                        cell.font = data_font
                        cell.alignment = center
                        cell.border = border

                ws.column_dimensions['A'].width = 14
                for c in range(2, len(header_row) + 1):
                    ws.column_dimensions[get_column_letter(c)].width = 8

                ws.row_dimensions[1].height = 20
                ws.row_dimensions[header_row_idx].height = 20

            # --- Лист «Методика» ---
            ws_m = wb.create_sheet("Методика")

            ws_m.column_dimensions['A'].width = 4
            for c in range(2, 12):
                ws_m.column_dimensions[get_column_letter(c)].width = 14

            row = 1

            ws_m.cell(row=row, column=2,
                      value="МЕТОДИКА ПЕРЕСЧЁТА МОЩНОСТЕЙ РАДИАТОРОВ LAGGARTT").font = title_font
            ws_m.merge_cells(start_row=row, start_column=2, end_row=row, end_column=10)
            row += 1

            ws_m.cell(row=row, column=2, value="Согласно ГОСТ 31311-2022").font = data_font
            ws_m.merge_cells(start_row=row, start_column=2, end_row=row, end_column=10)
            row += 2

            graph_str = f"{self._graph['supply']}/{self._graph['return']}/{self._graph['room']}"
            delta_str = f"{self._delta_t:.2f}".replace('.', ',')
            ws_m.cell(row=row, column=2,
                      value=f"Текущий температурный график: {graph_str} °С   (Δt = {delta_str} °С)").font = section_font
            ws_m.merge_cells(start_row=row, start_column=2, end_row=row, end_column=10)
            row += 2

            # Блок 1
            ws_m.cell(row=row, column=2, value="1. ЭКСПОНЕНТЫ n ПО ТИПАМ РАДИАТОРОВ").font = section_font
            ws_m.merge_cells(start_row=row, start_column=2, end_row=row, end_column=10)
            row += 1

            exp_text = (
                "Экспонент n определяет чувствительность мощности к изменению температурного напора Δt. "
                "Чем больше n, тем сильнее падает мощность при снижении Δt. "
                "Значения n одинаковы для всех высот (300–900 мм)."
            )
            ws_m.cell(row=row, column=2, value=exp_text).font = data_font
            ws_m.cell(row=row, column=2).alignment = left_wrap
            ws_m.merge_cells(start_row=row, start_column=2, end_row=row, end_column=10)
            ws_m.row_dimensions[row].height = 45
            row += 2

            exp_headers = ["Тип"] + VALID_TYPES
            for c_idx, val in enumerate(exp_headers, start=2):
                cell = ws_m.cell(row=row, column=c_idx, value=val)
                cell.font = header_font
                cell.alignment = center
                cell.border = border
                cell.fill = header_fill
            row += 1

            exp_row = ["n"]
            for rt in VALID_TYPES:
                n_val = EXPONENTS.get(500, {}).get(rt, 1.0)
                exp_row.append(f"{n_val:.2f}".replace('.', ','))
            for c_idx, val in enumerate(exp_row, start=2):
                cell = ws_m.cell(row=row, column=c_idx, value=val)
                cell.font = data_font
                cell.alignment = center
                cell.border = border
            row += 3

            # Блок 2
            ws_m.cell(row=row, column=2, value="2. ФОРМУЛА ПЕРЕСЧЁТА").font = section_font
            ws_m.merge_cells(start_row=row, start_column=2, end_row=row, end_column=10)
            row += 1

            formula_text = (
                "Мощность при текущем графике:\n"
                "    Q = Qном × (Δt / 70)ⁿ\n\n"
                "где Qном — паспортная мощность при 95/85/20 (Δt = 70 °С),\n"
                "    Δt — текущий температурный напор,\n"
                "    n — экспонент типа радиатора.\n\n"
                "Температурный напор:\n"
                "    Δt = (t подачи + t обратки) / 2 − t помещения\n\n"
                "При Δt = 70 коэффициент равен 1,0 — пересчёт не меняет паспортные значения."
            )
            ws_m.cell(row=row, column=2, value=formula_text).font = mono_font
            ws_m.cell(row=row, column=2).alignment = left_wrap
            ws_m.cell(row=row, column=2).fill = formula_fill
            ws_m.merge_cells(start_row=row, start_column=2, end_row=row, end_column=10)
            ws_m.row_dimensions[row].height = 140
            row += 3

            # Блок 3
            ws_m.cell(row=row, column=2, value="3. ПРИМЕР РАСЧЁТА ПРИ ТЕКУЩЕМ ГРАФИКЕ").font = section_font
            ws_m.merge_cells(start_row=row, start_column=2, end_row=row, end_column=10)
            row += 1

            qnom = 914
            try:
                df = self.app.data_provider.radiators_df
                mask = (
                    df['Наименование'].str.contains('/500/1000', regex=False, na=False) &
                    (df['RadiatorType'] == 10)
                )
                matches = df[mask]
                if not matches.empty:
                    qnom = int(round(float(matches.iloc[0].get('Мощность, Вт', 914))))
            except Exception:
                pass

            n = EXPONENTS.get(500, {}).get('10', 1.14)
            factor = (self._delta_t / 70.0) ** n if self._delta_t > 0 else 1.0
            q_recalc = qnom * factor if not PowerRecalculator.is_reference(self._delta_t) else qnom

            n_str = f"{n:.2f}".replace('.', ',')
            factor_str = f"{factor:.3f}".replace('.', ',')

            example_text = (
                f"Возьмём радиатор Тип 10, высота 500 мм, длина 1000 мм.\n\n"
                f"Паспортная мощность:  Qном = {qnom} Вт\n"
                f"Текущий график:       {graph_str}  →  Δt = {delta_str} °С\n"
                f"Экспонент (тип 10):   n = {n_str}\n"
                f"Коэффициент:          ({self._delta_t:.0f} / 70)^{n_str} = {factor_str}\n"
                f"Мощность:             Q = {qnom} × {factor_str} ≈ {int(round(q_recalc))} Вт"
            )
            ws_m.cell(row=row, column=2, value=example_text).font = mono_font
            ws_m.cell(row=row, column=2).alignment = left_wrap
            ws_m.cell(row=row, column=2).fill = formula_fill
            ws_m.merge_cells(start_row=row, start_column=2, end_row=row, end_column=10)
            ws_m.row_dimensions[row].height = 130
            row += 3

            # Блок 4
            ws_m.cell(row=row, column=2, value="4. ПОЯСНЕНИЕ К ТЕМПЕРАТУРНОМУ ГРАФИКУ").font = section_font
            ws_m.merge_cells(start_row=row, start_column=2, end_row=row, end_column=10)
            row += 1

            graph_text = (
                "Текущий график записывается тремя числами через слэш:\n"
                f"    {self._graph['supply']} / {self._graph['return']} / {self._graph['room']}\n\n"
                f"{self._graph['supply']} — температура подающей линии, °С\n"
                f"{self._graph['return']} — температура обратной линии, °С\n"
                f"{self._graph['room']} — температура воздуха в помещении, °С\n\n"
                "Типовые графики и соответствующие Δt:"
            )
            ws_m.cell(row=row, column=2, value=graph_text).font = data_font
            ws_m.cell(row=row, column=2).alignment = left_wrap
            ws_m.merge_cells(start_row=row, start_column=2, end_row=row, end_column=10)
            ws_m.row_dimensions[row].height = 110
            row += 2

            graph_headers = ["График", "Δt, °С"]
            for c_idx, val in enumerate(graph_headers, start=2):
                cell = ws_m.cell(row=row, column=c_idx, value=val)
                cell.font = header_font
                cell.alignment = center
                cell.border = border
                cell.fill = header_fill
            row += 1

            graph_rows = [
                ("95/85/20 (эталон)", "70,0"),
                ("90/70/20", "60,0"),
                ("85/60/20", "52,5"),
                ("75/65/20", "50,0"),
                ("55/45/20", "30,0"),
            ]
            for gr, dt_val in graph_rows:
                cell1 = ws_m.cell(row=row, column=2, value=gr)
                cell1.font = data_font
                cell1.alignment = center
                cell1.border = border

                cell2 = ws_m.cell(row=row, column=3, value=dt_val)
                cell2.font = data_font
                cell2.alignment = center
                cell2.border = border
                row += 1

            wb.save(str(file_path))
            logger.info(f"Таблицы мощностей сохранены: {file_path}")

            if platform.system() == "Windows":
                os.startfile(str(file_path))
            elif platform.system() == "Darwin":
                subprocess.call(["open", str(file_path)])
            else:
                subprocess.call(["xdg-open", str(file_path)])

        except Exception as e:
            logger.error(f"Ошибка экспорта мощностей: {e}")
            import traceback
            logger.error(traceback.format_exc())
            messagebox.showerror("Ошибка", f"Не удалось сохранить файл:\n{str(e)}")

    def _close(self) -> None:
        """Закрывает окно и отписывается от события масштабирования."""
        try:
            bus.unsubscribe(Events.FONT_SCALE_CHANGED, self._on_font_scale_changed)
        except Exception as e:
            logger.debug(f"Ошибка отписки от FONT_SCALE_CHANGED: {e}")

        if self._dialog:
            self._dialog.destroy()
            self._dialog = None
            self._tree = None
            self._table_frame = None
            self._canvas = None
            self._content_frame = None
            self._table_font = None
            self._table_font_bold = None