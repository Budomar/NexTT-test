"""
Запись спецификации в Excel-файл.
Создаёт два листа: Спецификация и Таблица соответствия (если есть данные).
"""

import os
import tempfile
from datetime import datetime
from typing import Optional, Tuple

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter

from nextt.logger import get_logger

logger = get_logger(__name__)


class ExcelWriter:
    """Записывает спецификацию в Excel."""

    # Стили
    HEADER_FONT = Font(name='Calibri', size=11, bold=True)
    DATA_FONT = Font(name='Calibri', size=11)
    BOLD_FONT = Font(name='Calibri', size=11, bold=True)
    ITALIC_FONT = Font(name='Calibri', size=11, italic=True, color="595959")

    ALIGN_CENTER = Alignment(horizontal='center', vertical='center')
    ALIGN_LEFT = Alignment(horizontal='left', vertical='center')

    THIN_BORDER = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin'),
    )

    MONEY_FORMAT = '#,##0.00'

    def __init__(self, data_provider=None):
        """
        Args:
            data_provider: экземпляр DataProvider для поиска мощности по артикулу
        """
        self.data_provider = data_provider

    def save(self, spec_data: pd.DataFrame, file_path: str,
             correspondence_data: Optional[pd.DataFrame] = None,
             program_name: str = "NexTT") -> None:
        """
        Сохраняет спецификацию в Excel.

        Args:
            spec_data: DataFrame с данными спецификации
            file_path: путь для сохранения
            correspondence_data: DataFrame с таблицей соответствия (опционально)
            program_name: название программы для подписи (берётся из заголовка окна)
        """
        wb = Workbook()
        ws = wb.active
        ws.title = "Спецификация"

        headers = [
            "№", "Артикул", "Наименование", "Мощность, Вт",
            "Цена, руб (с НДС)", "Скидка, %",
            "Цена со скидкой, руб (с НДС)", "Кол-во",
            "Сумма, руб (с НДС)"
        ]
        ws.append(headers)

        # Стили заголовков
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = self.HEADER_FONT
            cell.alignment = self.ALIGN_CENTER
            cell.border = self.THIN_BORDER

        # Данные
        for i, (_, row) in enumerate(spec_data.iterrows(), 2):
            is_bracket = "Кронштейн" in str(row.get("Наименование", ""))
            power_value = "" if is_bracket else row.get("Мощность, Вт", "")

            ws.append([
                i - 1,
                str(row.get("Артикул", "")),
                row.get("Наименование", ""),
                power_value,
                float(row.get("Цена, руб (с НДС)", 0)),
                float(row.get("Скидка, %", 0)),
                float(row.get("Цена со скидкой, руб (с НДС)", 0)),
                int(row.get("Кол-во", 0)),
                float(row.get("Сумма, руб (с НДС)", 0)),
            ])

            for col in range(1, len(headers) + 1):
                cell = ws.cell(row=i, column=col)
                cell.font = self.DATA_FONT
                cell.border = self.THIN_BORDER
                if col in [5, 7, 9]:
                    cell.number_format = self.MONEY_FORMAT
                    cell.alignment = self.ALIGN_CENTER
                elif col == 4:
                    cell.alignment = self.ALIGN_CENTER
                elif col in [1, 6, 8]:
                    cell.alignment = self.ALIGN_CENTER
                else:
                    cell.alignment = self.ALIGN_LEFT

        # Итого в спецификации
        total_row = len(spec_data) + 2
        total_sum = spec_data["Сумма, руб (с НДС)"].sum()

        total_qty_rad = sum(
            int(row["Кол-во"]) for _, row in spec_data.iterrows()
            if "Кронштейн" not in str(row.get("Наименование", ""))
        )
        total_qty_br = sum(
            int(row["Кол-во"]) for _, row in spec_data.iterrows()
            if "Кронштейн" in str(row.get("Наименование", ""))
        )

        ws.append(["Итого", "", "", "", "", "", "",
                   f"{total_qty_rad}/{total_qty_br}", total_sum])

        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=total_row, column=col)
            cell.font = self.BOLD_FONT
            cell.border = self.THIN_BORDER
            cell.alignment = self.ALIGN_CENTER
            if col in [5, 7, 9]:
                cell.number_format = self.MONEY_FORMAT

        # Расчёт суммарного веса и объёма радиаторов (без кронштейнов)
        total_weight, total_volume = self._calculate_total_weight_and_volume(spec_data)

        # Пустая строка перед весом и объёмом
        ws.append([])

        # Строка с весом
        ws.append([f"Суммарный вес радиаторов без учета упаковки и кронштейнов - {total_weight} кг."])
        ws.merge_cells(start_row=total_row + 2, start_column=1, end_row=total_row + 2, end_column=9)
        cell = ws.cell(row=total_row + 2, column=1)
        cell.font = self.DATA_FONT
        cell.alignment = self.ALIGN_LEFT

        # Строка с объёмом
        ws.append([f"Суммарный объем радиаторов без учета упаковки и кронштейнов - {total_volume} м3."])
        ws.merge_cells(start_row=total_row + 3, start_column=1, end_row=total_row + 3, end_column=9)
        cell = ws.cell(row=total_row + 3, column=1)
        cell.font = self.DATA_FONT
        cell.alignment = self.ALIGN_LEFT

        # Пустая строка перед подписью
        ws.append([])

        # Подпись с использованием переданного имени программы
        ws.append([f"Подготовлено в {program_name}"])
        ws.merge_cells(start_row=total_row + 5, start_column=1,
                       end_row=total_row + 5, end_column=9)
        cell = ws.cell(row=total_row + 5, column=1)
        cell.font = self.ITALIC_FONT
        cell.alignment = self.ALIGN_LEFT

        # Ширина столбцов
        widths = {'A': 5, 'B': 12, 'C': 50, 'D': 15, 'E': 20,
                  'F': 10, 'G': 40, 'H': 12, 'I': 20}
        for col_letter, width in widths.items():
            ws.column_dimensions[col_letter].width = width

        # ================================================================
        # Лист "Таблица соответствия" (если есть данные)
        # ================================================================
        if correspondence_data is not None and not correspondence_data.empty:
            # Удаляем существующий лист "Таблица соответствия", если он есть
            if "Таблица соответствия" in wb.sheetnames:
                std = wb["Таблица соответствия"]
                wb.remove(std)
            
            ws_corr = wb.create_sheet("Таблица соответствия")

            # Обрабатываем данные — добавляем мощность
            processed_data = []
            for _, row in correspondence_data.iterrows():
                meteor_art = str(row.get('Артикул LaggarTT', '')).strip()
                power_value = self._find_power_by_article(meteor_art)

                # Пытаемся найти количество в разных возможных столбцах
                qty_value = 0
                if 'Кол-во' in row:
                    qty_value = row.get('Кол-во', 0)
                elif 'Количество' in row:
                    qty_value = row.get('Количество', 0)
                else:
                    qty_value = row.get(list(row.index)[1], 0) if len(row.index) > 1 else 0
                
                # Пытаемся найти наименование в разных возможных столбцах
                original_name = ''
                if 'Наименование' in row:
                    original_name = row.get('Наименование', '')
                elif 'Оригинальное наименование' in row:
                    original_name = row.get('Оригинальное наименование', '')
                else:
                    original_name = row.get(list(row.index)[0], '') if len(row.index) > 0 else ''
                
                # Пытаемся найти наименование LaggarTT
                meteor_name = ''
                if 'Наименование LaggarTT' in row:
                    meteor_name = row.get('Наименование LaggarTT', '')
                else:
                    meteor_name = row.get(list(row.index)[2], '') if len(row.index) > 2 else ''

                processed_data.append({
                    'Оригинальное наименование': original_name,
                    'Количество': qty_value,
                    'Наименование LaggarTT': meteor_name,
                    'Артикул LaggarTT': meteor_art,
                    'Мощность, Вт': power_value
                })

            corr_headers = [
                'Оригинальное наименование',
                'Количество',
                'Наименование LaggarTT',
                'Артикул LaggarTT',
                'Мощность, Вт'
            ]
            ws_corr.append(corr_headers)

            # Стили заголовков листа соответствия
            for col in range(1, len(corr_headers) + 1):
                cell = ws_corr.cell(row=1, column=col)
                cell.font = self.HEADER_FONT
                cell.alignment = self.ALIGN_CENTER
                cell.border = self.THIN_BORDER

            # Данные листа соответствия
            total_qty_corr = 0

            for i, row_data in enumerate(processed_data, 2):
                ws_corr.append([
                    row_data['Оригинальное наименование'],
                    row_data['Количество'],
                    row_data['Наименование LaggarTT'],
                    row_data['Артикул LaggarTT'],
                    row_data['Мощность, Вт']
                ])

                art = str(row_data['Артикул LaggarTT']).strip()
                if art and art != '' and art != 'None':
                    try:
                        total_qty_corr += int(row_data['Количество'])
                    except (ValueError, TypeError):
                        pass

                for col in range(1, len(corr_headers) + 1):
                    cell = ws_corr.cell(row=i, column=col)
                    cell.font = self.DATA_FONT
                    cell.border = self.THIN_BORDER
                    if col in [2, 5]:
                        cell.alignment = self.ALIGN_CENTER
                    else:
                        cell.alignment = self.ALIGN_LEFT

            # Добавляем итоговую строку
            total_row_corr = len(processed_data) + 2
            ws_corr.append(["Итого", "", "", "", ""])
            ws_corr.cell(row=total_row_corr, column=2, value=total_qty_corr)

            RED_FONT = Font(name='Calibri', size=11, bold=True, color="FF0000")
            
            if total_qty_corr != total_qty_rad:
                cell = ws_corr.cell(row=total_row_corr, column=2)
                cell.font = RED_FONT
                cell.border = self.THIN_BORDER
                cell.alignment = self.ALIGN_CENTER
            else:
                cell = ws_corr.cell(row=total_row_corr, column=2)
                cell.font = self.BOLD_FONT
                cell.border = self.THIN_BORDER
                cell.alignment = self.ALIGN_CENTER

            for col in range(1, len(corr_headers) + 1):
                if col != 2:
                    cell = ws_corr.cell(row=total_row_corr, column=col)
                    cell.font = self.BOLD_FONT
                    cell.border = self.THIN_BORDER
                    if col in [5]:
                        cell.alignment = self.ALIGN_CENTER
                    else:
                        cell.alignment = self.ALIGN_LEFT

            ws_corr.merge_cells(start_row=total_row_corr, start_column=1, 
                                end_row=total_row_corr, end_column=1)

            # Автоподбор ширины столбцов
            for col_idx, column_name in enumerate(corr_headers, 1):
                max_length = len(str(column_name))
                column_letter = get_column_letter(col_idx)
                for row in ws_corr.iter_rows(min_col=col_idx, max_col=col_idx):
                    for cell in row:
                        try:
                            cell_length = len(str(cell.value))
                            if cell_length > max_length:
                                max_length = cell_length
                        except:
                            pass
                adjusted_width = (max_length + 2) * 1.2
                ws_corr.column_dimensions[column_letter].width = adjusted_width

        wb.save(file_path)
        logger.info(f"Спецификация сохранена: {file_path}")

    def _calculate_total_weight_and_volume(self, spec_data: pd.DataFrame) -> Tuple[float, float]:
        """
        Рассчитывает общий вес и объём радиаторов (без кронштейнов).

        Args:
            spec_data: DataFrame с данными спецификации

        Returns:
            (общий_вес_кг, общий_объём_м3) - оба числа округлены
        """
        total_weight = 0.0
        total_volume = 0.0

        if self.data_provider is None:
            return round(total_weight, 1), round(total_volume, 3)

        radiators_df = self.data_provider.radiators_df

        for _, row in spec_data.iterrows():
            # Пропускаем кронштейны
            if "Кронштейн" in str(row.get("Наименование", "")):
                continue

            art = str(row.get("Артикул", "")).strip()
            if not art:
                continue

            try:
                qty = int(row.get("Кол-во", 0))
                if qty <= 0:
                    continue
            except (ValueError, TypeError):
                continue

            # Ищем радиатор по артикулу
            mask = radiators_df['Артикул'].astype(str).str.strip() == art
            if mask.any():
                product = radiators_df[mask].iloc[0]
                total_weight += float(product.get('Вес, кг', 0)) * qty
                total_volume += float(product.get('Объем, м3', 0)) * qty

        return round(total_weight, 1), round(total_volume, 3)

    def _find_power_by_article(self, article: str):
        """Находит мощность радиатора по артикулу."""
        if not article or not self.data_provider:
            return ""
        found = self.data_provider.find_by_article(article)
        if found is not None:
            power = found.get('Мощность, Вт', '')
            try:
                if pd.isna(power):
                    return ""
                return float(power)
            except (ValueError, TypeError):
                return str(power) if not pd.isna(power) else ""
        return ""