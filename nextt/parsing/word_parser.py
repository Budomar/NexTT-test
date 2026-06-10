"""
Парсер Word-файлов для извлечения таблиц.
"""

import os
import pandas as pd

from nextt.logger import get_logger

logger = get_logger(__name__)

try:
    from docx import Document
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False
    logger.warning("python-docx не установлен. Установите: pip install python-docx")


class WordParser:
    """Читает таблицы из Word-файлов."""

    @classmethod
    def parse(cls, file_path: str) -> pd.DataFrame:
        """
        Читает Word-файл и возвращает DataFrame со всеми таблицами.

        Args:
            file_path: путь к .docx файлу

        Returns:
            DataFrame с данными (все значения как строки)

        Raises:
            ImportError: если python-docx не установлен
            ValueError: если формат не поддерживается
        """
        if not HAS_DOCX:
            raise ImportError(
                "Для чтения Word-файлов требуется python-docx. "
                "Установите: pip install python-docx"
            )

        ext = os.path.splitext(file_path)[1].lower()
        if ext not in ['.docx']:
            raise ValueError(f"Неподдерживаемый формат: {ext}. Ожидается .docx")

        logger.info(f"Чтение Word-файла: {os.path.basename(file_path)}")

        try:
            doc = Document(file_path)
        except Exception as e:
            raise RuntimeError(f"Ошибка чтения Word-файла: {e}")

        all_rows = []

        for table_idx, table in enumerate(doc.tables):
            for row_idx, row in enumerate(table.rows):
                cells = []
                for cell in row.cells:
                    # Объединяем текст из параграфов ячейки
                    text = ' '.join(p.text for p in cell.paragraphs).strip()
                    cells.append(text)
                
                # Пропускаем полностью пустые строки
                if any(cells):
                    all_rows.append(cells)

        if not all_rows:
            logger.warning("Таблицы не найдены или пусты")
            return pd.DataFrame()

        # Выравниваем количество столбцов (разные таблицы могут иметь разное число столбцов)
        max_cols = max(len(row) for row in all_rows)
        for row in all_rows:
            while len(row) < max_cols:
                row.append('')

        df = pd.DataFrame(all_rows)
        df = df.fillna('')
        logger.info(f"Извлечено: {df.shape[0]} строк, {df.shape[1]} столбцов")
        return df